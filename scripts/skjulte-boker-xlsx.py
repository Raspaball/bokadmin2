"""Lager scripts/out/skjulte-boker-3280-detaljert.xlsx fra scripts/out/skjulte-detaljer.json.
Fane 1: én rad per bok. Fane 2: sammendrag. Fane 3: antakelser. Bare lokalt (live-data, ikke i git)."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

out = Path(__file__).parent / "out"
data = json.loads((out / "skjulte-detaljer.json").read_text(encoding="utf-8"))
rows, s = data["rows"], data["sammendrag"]

cols = [
    ("Tittel", "tittel"), ("ISBN", "isbn"), ("Forlag", "forlag"), ("Pris i dag", "pris_i_dag"), ("Pris fra Bokbasen", "pris_bokbasen"),
    ("Pris 0 / mangler gyldig pris", "pris_null_eller_mangler"), ("Utgivelsesdato", "utgivelsesdato"), ("Kommende (kode 10–12)", "kommende"),
    ("ONIX-kode", "onix_kode"), ("Status i dag", "status_i_dag"), ("Ny status etter regelen", "ny_status"),
    ("Fysisk bok", "fysisk_bok"), ("Produktform (ONIX)", "produktform"), ("Finnes hos Bokbasen", "i_bokbasen"),
    ("Duplikat-ISBN", "duplikat"), ("Beskyttet", "beskyttet"), ("Produkttype i dag", "produkttype_i_dag"),
    ("Kanaler i dag", "kanaler_i_dag"), ("Antall kanaler", "antall_kanaler"),
    ("Ville blitt synlig i nettbutikken (regelen i dag)", "synlig_i_nettbutikken"),
    ("Hadde blitt synlig hvis produkttype var «Bok» (ANTAKELSE)", "publiseres_hvis_produkttype_bok"),
    ("Anbefaling", "anbefaling"), ("Opprettet", "opprettet"),
]
wb = Workbook()
ws = wb.active
ws.title = "Bøker"
head = PatternFill("solid", fgColor="1F3A5F")
ws.append([c[0] for c in cols])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = head; c.alignment = Alignment(wrap_text=True, vertical="top")
for r in rows:
    ws.append([r.get(k) for _, k in cols])
widths = {"Tittel": 42, "Forlag": 22, "Kanaler i dag": 40, "Ville blitt synlig i nettbutikken (regelen i dag)": 46, "Anbefaling": 46}
for i, (n, _) in enumerate(cols, 1):
    ws.column_dimensions[get_column_letter(i)].width = widths.get(n, 16)
ws.row_dimensions[1].height = 62
ws.freeze_panes = "B2"
ws.auto_filter.ref = ws.dimensions
fills = {"publiser": "C6EFCE", "hold utenfor": "F4CCCC", "sjekk manuelt": "FFF2CC"}
ai = [k for _, k in cols].index("anbefaling") + 1
for row in ws.iter_rows(min_row=2, min_col=ai, max_col=ai):
    for c in row:
        for key, color in fills.items():
            if str(c.value or "").startswith(key):
                c.fill = PatternFill("solid", fgColor=color)

ss = wb.create_sheet("Sammendrag")
def put(title, items=None, note=None):
    ss.append([title]); ss.cell(ss.max_row, 1).font = Font(bold=True)
    if note: ss.append([note])
    for k, v in (items or []): ss.append([str(k), v])
    ss.append([])
put(f"Skjulte aktive bøker (generert {s['generert']}): {s['antall']}")
put("Fra 3 280 til de som ville blitt publisert (kumulativt, i samme rekkefølge som jobbene sjekker)", s["trinn"])
put("Anbefaling (bare sjekkmodus-regelen)", sorted(s["anbefaling"].items(), key=lambda x: -x[1]))
put("Første grunn til at de ikke ville blitt publisert", sorted(s["hvorfor"].items(), key=lambda x: -x[1]))
put("ONIX-kode", sorted(s["onixKode"].items(), key=lambda x: -x[1]))
put("Ny status etter regelen", sorted(s["nyStatus"].items(), key=lambda x: -x[1]))
put("Fysisk bok", sorted(s["fysisk"].items(), key=lambda x: -x[1]))
put("Finnes hos Bokbasen", sorted(s["iBokbasen"].items(), key=lambda x: -x[1]))
put("Kanaler i dag", sorted(s["kanalerIdag"].items(), key=lambda x: -x[1]))
put("Produkttype i dag", list(s["produkttypeIdag"].items()))
put("Forlag (de 8 største)", s["topForlag"])
put("Opprettet (måned)", sorted(s["opprettet"].items()))
put("Annet", [("Kommende bøker (kode 10–12)", s["kommende"]), ("Pris 0 eller mangler gyldig pris", s["prisNullEllerMangler"]),
              ("Publiseres av regelen i dag", s["publiseresIDag"]), ("Ville publiseres hvis produkttype var «Bok» (antakelse)", s["publiseresHvisBok"])])
ss.column_dimensions["A"].width = 78; ss.column_dimensions["B"].width = 14

aa = wb.create_sheet("Antakelser")
for line in [
    "Dette er en kartlegging i sjekkmodus. Ingenting er publisert eller endret i butikken.",
    "«Ville blitt synlig» følger regelen channelsToPublish() slik book-update og availability-check bruker den i dag.",
    "ANTAKELSE: kolonnen «hvis produkttype var Bok» viser hva regelen ville gjort etter at bokdata-jobben har satt produkttypen til «Bok». Det er ikke gjort i live.",
    "ANTAKELSE: ny status følger availabilityRule() på ONIX-koden fra Bokbasen i dag. Bøker som blir utkast eller arkivert publiseres ikke.",
    "Anbefaling «publiser» = regelen publiserer i dag. «sjekk manuelt» = bare produkttypen (eller pris) stopper. «hold utenfor» = regelen publiserer ikke uansett produkttype.",
    "Eirik avgjør om bøkene skal synliggjøres (åpne spørsmål 1 og 5). Anbefalingen er ikke et vedtak.",
    "Lista inneholder live-data (titler, ISBN). Den ligger bare lokalt og skal ikke i git.",
]:
    aa.append([line])
aa.column_dimensions["A"].width = 150
wb.save(out / "skjulte-boker-3280-detaljert.xlsx")
print("Lagret", out / "skjulte-boker-3280-detaljert.xlsx", len(rows), "rader")
